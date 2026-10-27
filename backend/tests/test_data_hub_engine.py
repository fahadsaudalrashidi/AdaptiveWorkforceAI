from __future__ import annotations

import io
import json
import unittest
from datetime import timedelta

import pandas as pd
from openpyxl import load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.api.data_hub_routes import create_data_hub_router
from backend.core.data_hub_engine import (
    apply_cleaning,
    build_summary_pdf,
    build_summary_xlsx,
    _clean_numeric,
    forecast_frame,
    profile_frame,
    workforce_scenario_frame,
)


class DataHubProfileTests(unittest.TestCase):
    def test_profile_returns_schema_aggregates_and_withholds_sensitive_statistics(self):
        frame = pd.DataFrame({
            "Department": ["Operations"] * 9 + ["Finance"],
            "Monthly Cost": [100, 101, 98, 99, 102, 100, 103, 97, 99, 1000],
            "Employee Name": [f"Person {i}" for i in range(10)],
        })

        result = profile_frame(frame, filename="costs.csv", sheet="Data")

        cost = next(column for column in result["columns"] if column["name"] == "Monthly Cost")
        name = next(column for column in result["columns"] if column["name"] == "Employee Name")
        self.assertEqual(result["source_rows"], 10)
        self.assertEqual(cost["outlier_count"], 1)
        self.assertEqual(name["sensitive"], True)
        self.assertNotIn("numeric_summary", name)
        self.assertFalse(any("Person" in str(value) for value in result.values()))
        self.assertEqual(result["anomalies"][0]["method"], "Tukey IQR fence (1.5 × IQR)")

    def test_cleaning_is_opt_in_and_discloses_changes(self):
        frame = pd.DataFrame({" Team ": [" Ops ", "Ops", None], "Empty": [None, None, None]})
        untouched, applied_none, counts_none = apply_cleaning(frame)
        self.assertEqual(len(untouched), 3)
        self.assertEqual(applied_none, [])

        cleaned, applied, counts = apply_cleaning(
            frame, ["trim_text", "remove_blank_rows", "remove_blank_columns", "remove_exact_duplicates"]
        )
        self.assertEqual(counts["source_rows"], 3)
        self.assertEqual(counts["analyzed_rows"], 1)
        self.assertEqual(counts["source_columns"], 2)
        self.assertEqual(counts["analyzed_columns"], 1)
        self.assertEqual(cleaned.columns.tolist(), [" Team "])
        self.assertEqual(cleaned.iloc[0, 0], "Ops")
        self.assertEqual({entry["id"] for entry in applied}, {
            "trim_text", "remove_blank_rows", "remove_blank_columns", "remove_exact_duplicates"
        })

    def test_expiry_flags_use_the_saudi_local_calendar_and_aggregate_counts(self):
        today = pd.Timestamp(pd.Timestamp.now(tz="Asia/Riyadh").date())
        frame = pd.DataFrame({
            "Iqama Expiry": [today - timedelta(days=1), today + timedelta(days=5),
                             today + timedelta(days=45), pd.NaT],
        })
        profile = profile_frame(frame, filename="visas.csv", sheet="Data")
        expiry = next(anomaly for anomaly in profile["anomalies"] if anomaly["kind"] == "compliance")
        self.assertEqual(expiry["expired"], 1)
        self.assertEqual(expiry["due_30_days"], 1)
        self.assertEqual(expiry["evaluated"], 3)

    def test_temporal_cost_spikes_compare_against_rolling_local_baseline(self):
        frame = pd.DataFrame({
            "Month": pd.date_range("2025-01-01", periods=12, freq="MS"),
            "Operating Cost": [100] * 11 + [220],
        })
        profile = profile_frame(frame, filename="costs.csv", sheet="Data")
        spikes = [
            anomaly for anomaly in profile["anomalies"]
            if anomaly["kind"] == "temporal_cost_spike"
        ]
        self.assertEqual(len(spikes), 1)
        self.assertEqual(spikes[0]["periods"], ["2025-12"])
        self.assertEqual(spikes[0]["count"], 1)
        self.assertIn("MAD", spikes[0]["method"])
        self.assertEqual(spikes[0]["seasonal_pattern_periods"], [])
        self.assertIn("not confirmed", spikes[0]["seasonality_context"])

    def test_seasonal_context_is_additive_and_requires_three_annual_cycles(self):
        values = [100] * 36
        values[11] = 220
        values[23] = 230
        values[35] = 240
        frame = pd.DataFrame({
            "Month": pd.date_range("2024-01-01", periods=36, freq="MS"),
            "Operating Cost": values,
        })
        profile = profile_frame(frame, filename="seasonal-costs.csv", sheet="Data")
        spikes = [
            anomaly for anomaly in profile["anomalies"]
            if anomaly["kind"] == "temporal_cost_spike"
        ]
        self.assertEqual(len(spikes), 1)
        self.assertEqual(spikes[0]["count"], 3)
        self.assertEqual(spikes[0]["periods"], ["2024-12", "2025-12", "2026-12"])
        self.assertEqual(spikes[0]["seasonal_pattern_periods"], ["2026-12"])
        self.assertEqual(spikes[0]["severity"], "moderate")

    def test_monthly_forecast_has_observed_and_intervals(self):
        frame = pd.DataFrame({
            "Month": pd.date_range("2025-01-01", periods=12, freq="MS"),
            "Operating Budget": [100 + i * 10 for i in range(12)],
        })

        result = forecast_frame(
            frame,
            filename="budget.csv",
            sheet="Data",
            date_column="Month",
            metric_column="Operating Budget",
            aggregation="sum",
            horizon=3,
        )

        self.assertEqual(result["observation_count"], 12)
        self.assertEqual(result["forecast"][0]["period"], "2026-01")
        self.assertAlmostEqual(result["forecast"][0]["estimate"], 220, places=5)
        self.assertEqual(result["selected_model"], "linear_trend")
        self.assertLess(result["forecast"][0]["lower"], result["forecast"][0]["estimate"])
        self.assertGreater(result["forecast"][0]["upper"], result["forecast"][0]["estimate"])
        self.assertTrue(result["limitations"])

    def test_forecast_selects_seasonal_baseline_when_backtesting_supports_it(self):
        values = [20 + (month % 12) * 5 for month in range(30)]
        frame = pd.DataFrame({
            "Month": pd.date_range("2023-01-01", periods=30, freq="MS"),
            "Demand": values,
        })
        result = forecast_frame(
            frame,
            filename="demand.csv",
            sheet="Data",
            date_column="Month",
            metric_column="Demand",
            aggregation="mean",
            horizon=2,
        )
        self.assertEqual(result["selected_model"], "seasonal_naive")
        self.assertEqual(result["validation_periods"], 18)
        self.assertAlmostEqual(result["forecast"][0]["estimate"], values[-12], places=5)

    def test_forecast_rejects_sensitive_measure_and_short_history(self):
        frame = pd.DataFrame({
            "Month": pd.date_range("2025-01-01", periods=12, freq="MS"),
            "Salary": [100 + i for i in range(12)],
        })
        with self.assertRaisesRegex(ValueError, "personal or post-event"):
            forecast_frame(frame, filename="hr.csv", sheet="Data", date_column="Month", metric_column="Salary")

        short = frame.iloc[:7].copy()
        short["Operating Budget"] = range(7)
        with self.assertRaisesRegex(ValueError, "At least 8"):
            forecast_frame(short, filename="budget.csv", sheet="Data", date_column="Month", metric_column="Operating Budget")

    def test_workbook_export_contains_no_source_records(self):
        profile = profile_frame(
            pd.DataFrame({"Department": ["Secret Division", "Secret Division"], "Cost": [10, 20]}),
            filename="internal.csv",
            sheet="Data",
        )
        content = build_summary_xlsx([profile])
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        values = [cell for row in workbook["Data Profile"].iter_rows(values_only=True) for cell in row]
        self.assertIn("Cost", values)
        self.assertNotIn("Secret Division", values)
        workbook.close()

    def test_arabic_workbook_export_localizes_tabs_and_supports_rtl(self):
        profile = profile_frame(
            pd.DataFrame({"Department": ["Operations"] * 10, "Budget": list(range(10))}),
            filename="budget.csv",
            sheet="Data",
        )
        workbook = load_workbook(io.BytesIO(build_summary_xlsx([profile], language="ar")), read_only=False)
        self.assertEqual(workbook.sheetnames[0], "نظرة عامة")
        self.assertEqual(workbook["نظرة عامة"]["A1"].value, "THRIVE — ملخص مركز البيانات")
        self.assertTrue(workbook["نظرة عامة"].sheet_view.rightToLeft)
        self.assertEqual(workbook["ملف البيانات"]["A1"].value, "الملف")
        workbook.close()

    def test_pdf_export_is_a_valid_document_in_both_languages(self):
        profile = profile_frame(
            pd.DataFrame({"Department": ["Operations"] * 10, "Budget": list(range(10))}),
            filename="budget.csv",
            sheet="Data",
        )
        for language in ("en", "ar"):
            pdf = build_summary_pdf([profile], language)
            self.assertTrue(pdf.startswith(b"%PDF-"))
            self.assertGreater(len(pdf), 1000)

    def test_numeric_parser_does_not_extract_digits_from_labels(self):
        parsed = _clean_numeric(pd.Series(["Person 2026", "SAR 1,234.50", "٣٬٥٠٠٫٥٠", "(25.5)"]))
        self.assertTrue(pd.isna(parsed.iloc[0]))
        self.assertAlmostEqual(parsed.iloc[1], 1234.5)
        self.assertAlmostEqual(parsed.iloc[2], 3500.5)
        self.assertAlmostEqual(parsed.iloc[3], -25.5)

    def test_authenticated_data_hub_inspection_and_forecast_routes(self):
        router = create_data_hub_router()
        app = FastAPI()
        app.dependency_overrides[router.dependencies[0].dependency] = lambda: None
        app.include_router(router)
        client = TestClient(app)
        csv = (
            "Month,Operating Budget,Employee Name\n"
            "2025-01-01,100,Private One\n"
            "2025-02-01,110,Private Two\n"
            "2025-03-01,120,Private Three\n"
            "2025-04-01,130,Private Four\n"
            "2025-05-01,140,Private Five\n"
            "2025-06-01,150,Private Six\n"
            "2025-07-01,160,Private Seven\n"
            "2025-08-01,170,Private Eight\n"
        ).encode()
        inspect_response = client.post(
            "/api/local/data-hub/inspect",
            files=[("files", ("budget.csv", csv, "text/csv"))],
            data={"cleaning_options": "[]"},
        )
        self.assertEqual(inspect_response.status_code, 200, inspect_response.text)
        payload = inspect_response.json()
        self.assertTrue(payload["privacy"]["persisted"] is False)
        self.assertNotIn("Private One", inspect_response.text)
        self.assertEqual(payload["datasets"][0]["date_columns"], ["Month"])

        forecast_response = client.post(
            "/api/local/data-hub/forecast",
            files={"file": ("budget.csv", csv, "text/csv")},
            data={
                "sheet": "Data",
                "date_column": "Month",
                "metric_column": "Operating Budget",
                "aggregation": "sum",
                "horizon": "2",
            },
        )
        self.assertEqual(forecast_response.status_code, 200, forecast_response.text)
        self.assertEqual(forecast_response.json()["observation_count"], 8)

        export_response = client.post(
            "/api/local/data-hub/export.xlsx",
            files=[("files", ("budget.csv", csv, "text/csv"))],
            data={"language": "ar", "cleaning_options": "[]"},
        )
        self.assertEqual(export_response.status_code, 200, export_response.text)
        exported = load_workbook(io.BytesIO(export_response.content), read_only=True, data_only=True)
        self.assertIn("نظرة عامة", exported.sheetnames)
        exported.close()

    def test_data_hub_routes_require_authentication_by_default(self):
        app = FastAPI()
        app.include_router(create_data_hub_router())
        response = TestClient(app).post(
            "/api/local/data-hub/inspect",
            files=[("files", ("budget.csv", b"Month,Cost\n2025-01,10\n", "text/csv"))],
        )
        self.assertEqual(response.status_code, 401)


class WorkforceScenarioTests(unittest.TestCase):
    def setUp(self):
        self.frame = pd.DataFrame({
            "Department": ["Operations", "Finance"],
            "Headcount": [10, 5],
            "Annual Payroll": [100_000, 50_000],
            "Overtime Cost": [5_000, 2_500],
        })
        self.scenarios = [{
            "id": "growth",
            "name": "Growth",
            "headcount_change_pct": 10,
            "payroll_growth_pct": 5,
            "overtime_reduction_pct": 20,
        }]

    def build(self, frame=None, scenarios=None):
        return workforce_scenario_frame(
            frame if frame is not None else self.frame,
            filename="workforce.csv",
            sheet="Data",
            department_column="Department",
            headcount_column="Headcount",
            payroll_column="Annual Payroll",
            overtime_column="Overtime Cost",
            scenarios=scenarios if scenarios is not None else self.scenarios,
        )

    def test_scenarios_compare_workforce_and_cost_with_explicit_assumptions(self):
        result = self.build()
        self.assertEqual(result["baseline"]["headcount"], 15)
        self.assertEqual(result["baseline"]["labor_cost"], 157500)
        self.assertEqual(result["scenarios"][0]["totals"]["headcount"], 16.5)
        self.assertEqual(result["scenarios"][0]["totals"]["payroll"], 173250)
        self.assertEqual(result["scenarios"][0]["totals"]["overtime"], 6600)
        self.assertEqual(result["scenarios"][0]["totals"]["labor_cost"], 179850)
        self.assertEqual(len(result["baseline"]["groups"]), 2)
        self.assertFalse(result["privacy"]["persisted"])
        self.assertFalse(result["privacy"]["source_rows_returned"])

    def test_small_department_details_are_withheld(self):
        result = self.build(frame=pd.DataFrame({
            "Department": ["Operations", "Finance"],
            "Headcount": [10, 4],
            "Annual Payroll": [100_000, 40_000],
            "Overtime Cost": [5_000, 2_000],
        }))
        self.assertTrue(result["baseline"]["group_details_withheld"])
        self.assertEqual(result["baseline"]["groups"], [])
        self.assertEqual(result["baseline"]["headcount"], 14)
        self.assertTrue(any("fewer than five" in item for item in result["limitations"]))

    def test_department_details_are_available_at_the_five_person_threshold(self):
        result = self.build(frame=pd.DataFrame({
            "Department": ["Operations", "Finance"],
            "Headcount": [10, 5],
            "Annual Payroll": [100_000, 50_000],
            "Overtime Cost": [5_000, 2_500],
        }))
        self.assertFalse(result["baseline"]["group_details_withheld"])
        self.assertEqual(len(result["baseline"]["groups"]), 2)

    def test_totals_only_comparison_does_not_require_a_department_dimension(self):
        result = workforce_scenario_frame(
            self.frame,
            filename="workforce.csv",
            sheet="Data",
            department_column="",
            headcount_column="Headcount",
            payroll_column="Annual Payroll",
            overtime_column=None,
            scenarios=self.scenarios,
        )
        self.assertEqual(result["baseline"]["headcount"], 15)
        self.assertEqual(result["baseline"]["labor_cost"], 150000)
        self.assertEqual(result["baseline"]["groups"], [])
        self.assertFalse(result["baseline"]["group_details_withheld"])

    def test_sensitive_columns_and_out_of_range_assumptions_are_rejected(self):
        sensitive = self.frame.assign(Salary=[10_000, 10_000])
        with self.assertRaisesRegex(ValueError, "personal or post-event"):
            workforce_scenario_frame(
                sensitive,
                filename="workforce.csv",
                sheet="Data",
                department_column="Department",
                headcount_column="Headcount",
                payroll_column="Salary",
                overtime_column=None,
                scenarios=self.scenarios,
            )
        invalid = [{**self.scenarios[0], "headcount_change_pct": 201}]
        with self.assertRaisesRegex(ValueError, "between -100% and 200%"):
            self.build(scenarios=invalid)

    def test_negative_costs_are_rejected(self):
        invalid = self.frame.copy()
        invalid.loc[0, "Annual Payroll"] = -1
        with self.assertRaisesRegex(ValueError, "must be non-negative"):
            self.build(frame=invalid)

    def test_scenario_endpoint_returns_aggregates_and_requires_authentication(self):
        router = create_data_hub_router()
        app = FastAPI()
        app.dependency_overrides[router.dependencies[0].dependency] = lambda: None
        app.include_router(router)
        client = TestClient(app)
        csv = (
            "Department,Headcount,Annual Payroll,Overtime Cost\n"
            "Operations,10,100000,5000\n"
            "Finance,5,50000,2500\n"
        ).encode()
        response = client.post(
            "/api/local/data-hub/scenario",
            files={"file": ("workforce.csv", csv, "text/csv")},
            data={
                "sheet": "Data",
                "department_column": "Department",
                "headcount_column": "Headcount",
                "payroll_column": "Annual Payroll",
                "overtime_column": "Overtime Cost",
                "scenarios": json.dumps(self.scenarios),
                "language": "en",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["baseline"]["headcount"], 15)
        self.assertFalse(payload["privacy"]["persisted"])
        self.assertNotIn("Operations,10", response.text)

        protected = FastAPI()
        protected.include_router(create_data_hub_router())
        unauthorized = TestClient(protected).post(
            "/api/local/data-hub/scenario",
            files={"file": ("workforce.csv", csv, "text/csv")},
            data={
                "department_column": "Department",
                "headcount_column": "Headcount",
                "payroll_column": "Annual Payroll",
                "scenarios": json.dumps(self.scenarios),
            },
        )
        self.assertEqual(unauthorized.status_code, 401)


if __name__ == "__main__":
    unittest.main()
