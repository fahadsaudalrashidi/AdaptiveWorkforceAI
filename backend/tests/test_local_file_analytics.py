from __future__ import annotations

import io
import unittest
from unittest.mock import patch

import pandas as pd

from backend.core.ai_business_analyst_engine import answer_question
from backend.core.excel_ingestion_engine import ExcelIngestionEngine
from backend.core.local_question_engine import analyze_local_question, resolve_local_follow_up


def _simple_searchable_pdf(text: str | list[str]) -> bytes:
    """Make a tiny valid PDF fixture without adding a PDF authoring dependency."""
    lines = [text] if isinstance(text, str) else text
    commands = ["BT /F1 12 Tf 72 720 Td"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"({escaped}) Tj 0 -20 Td")
    content = (" ".join(commands) + " ET").encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode())
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(pdf)


class LocalFileAnalyticsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = ExcelIngestionEngine()

    def test_nested_json_is_flattened_for_analysis(self) -> None:
        raw = b'{"records":[{"department":"Sales","net":{"revenue":125}},{"department":"Ops","net":{"revenue":75}}]}'
        result = self.engine.ingest(raw, "revenue.json")
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.row_count, 2)
        self.assertIn("net.revenue", result.df.columns)

    def test_workbook_returns_each_nonempty_sheet(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        first = workbook.active
        first.title = "North"
        first.append(["Department", "Sales"])
        first.append(["Retail", 10])
        second = workbook.create_sheet("South")
        second.append(["Department", "Sales"])
        second.append(["Online", 20])
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = self.engine.ingest_all(buffer.getvalue(), "regional.xlsx")
        self.assertEqual([name for name, item in result if item.success], ["North", "South"])
        self.assertEqual([item.df.iloc[0]["Sales"] for _, item in result], [10, 20])

    def test_searchable_pdf_becomes_grounded_extractable_text(self) -> None:
        result = self.engine.ingest(_simple_searchable_pdf("Contract renewal date is 2026-12-01."), "terms.pdf")
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.df.iloc[0]["PDF page"], 1)
        answer = analyze_local_question("What does the contract say about renewal?", {"terms.pdf": result.df})
        self.assertIsNotNone(answer)
        self.assertEqual(answer.operation, "document_search")
        self.assertIn("2026-12-01", answer.detail)

    def test_grouped_sum_is_deterministic_and_integrated_in_chat_answer(self) -> None:
        frame = pd.DataFrame({"Department": ["Retail", "Retail", "Online"], "Net Sales": [12.5, 7.5, 30.0]})
        answer = answer_question("Show total net sales by department", {}, use_llm=False, uploaded_frames={"sales.csv": frame})
        self.assertEqual(answer.mode, "tabular")
        self.assertIn("Retail: 20.00", answer.key_evidence[0].value)
        self.assertIn("Online: 30.00", answer.key_evidence[0].value)

    def test_sensitive_column_is_not_used_even_when_other_numbers_exist(self) -> None:
        frame = pd.DataFrame({"Salary": [100000, 120000], "Net Sales": [5, 8]})
        answer = analyze_local_question("What is the total salary?", {"payroll.csv": frame})
        self.assertEqual(answer.operation, "clarify")
        self.assertIn("personal-data column", answer.text)
        self.assertNotIn("220,000", answer.detail)

    def test_ambiguous_measures_ask_for_clarification(self) -> None:
        frame = pd.DataFrame({"Net Sales": [10, 20], "Gross Sales": [12, 22]})
        answer = analyze_local_question("What is the total sales?", {"sales.csv": frame})
        self.assertEqual(answer.operation, "clarify")
        self.assertIn("more than one numeric column", answer.text)

    def test_count_works_without_a_numeric_column(self) -> None:
        frame = pd.DataFrame({"Department": ["Retail", "Retail", "Online"]})
        answer = analyze_local_question("Count records by department", {"teams.csv": frame})
        self.assertEqual(answer.operation, "count_by")
        self.assertIn("Retail: 2 rows", answer.detail)

    def test_unknown_measure_does_not_substitute_the_only_numeric_column(self) -> None:
        frame = pd.DataFrame({"Sales": [10, 20]})
        answer = analyze_local_question("What is total profit?", {"sales.csv": frame})
        self.assertEqual(answer.operation, "clarify")

    def test_natural_category_filter_is_applied_to_sum(self) -> None:
        frame = pd.DataFrame({"Sales": [10, 20], "Department": ["Retail", "Online"]})
        answer = analyze_local_question("Total sales only for Retail", {"sales.csv": frame})
        self.assertEqual(answer.operation, "total")
        self.assertIn("10.00", answer.detail)
        self.assertEqual(answer.query_audit["cohorts"][0]["value"], 10)
        self.assertEqual(answer.query_audit["cohorts"][0]["matched_rows"], 1)

    def test_governed_upload_preserves_json_and_pdf_formats(self) -> None:
        from backend.api.local_analysis_routes import _ingest_analyst_upload

        for filename, raw in [
            ("revenue.json", b'[{"Revenue":125},{"Revenue":75}]'),
            ("terms.pdf", _simple_searchable_pdf("Contract renewal is due in December.")),
        ]:
            frames, _, details = _ingest_analyst_upload(raw, filename)
            self.assertTrue(frames, details)
            self.assertFalse(any(detail.get("error") for detail in details), details)

    def test_borderless_pdf_table_can_be_calculated_locally(self) -> None:
        raw = _simple_searchable_pdf([
            "Department      Sales",
            "Retail          12.5",
            "Online          30.0",
            "Retail          7.5",
        ])
        result = self.engine.ingest(raw, "sales.pdf")
        self.assertTrue(result.success, result.error)
        self.assertIn("Sales", result.df.columns)
        answer = analyze_local_question("Total Sales", {"sales.pdf": result.df})
        self.assertIn("50.00", answer.detail)

    def test_json_collections_remain_separate_sources(self) -> None:
        from backend.api.local_analysis_routes import _ingest_analyst_upload

        raw = b'{"north":[{"Sales":10}],"south":[{"Sales":20}]}'
        frames, _, details = _ingest_analyst_upload(raw, "regions.json")
        self.assertEqual(len(frames), 2, details)
        answer = analyze_local_question("Total Sales", frames)
        self.assertEqual(answer.operation, "clarify")
        answer = analyze_local_question("Total Sales in north", frames)
        self.assertIn("10.00", answer.detail)

    def test_workbook_limit_is_explicit_not_a_silent_single_sheet_fallback(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        for index in range(50):
            workbook.create_sheet(f"Extra {index}")
        buffer = io.BytesIO()
        workbook.save(buffer)
        result = self.engine.ingest_all(buffer.getvalue(), "large.xlsx")
        self.assertFalse(result[0][1].success)
        self.assertIn("more than 50", result[0][1].error)

    def test_pdf_table_upload_retains_document_text(self) -> None:
        from backend.api.local_analysis_routes import _ingest_analyst_upload

        raw = _simple_searchable_pdf([
            "Department      Sales", "Retail          12.5",
            "Online          30.0", "Retail          7.5",
        ])
        frames, _, details = _ingest_analyst_upload(raw, "sales.pdf")
        self.assertEqual(len(frames), 2, details)
        answer = analyze_local_question("Summarize this document", frames)
        self.assertEqual(answer.operation, "document_search")
        self.assertIn("Retail", answer.detail)

    def test_chat_api_defaults_local_even_when_cloud_is_configured(self) -> None:
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from backend.api.local_analysis_routes import create_local_analysis_router
        from backend.core.auth_session import AWIRole
        from backend.core.cloud_llm_client import CloudModeStatus
        from backend.core.local_auth import LocalUser
        from backend.core.security_middleware import get_current_user

        app = FastAPI()
        app.include_router(create_local_analysis_router())
        user = LocalUser("test-user", "test", "Test", "", AWIRole.SYSTEM_ADMIN, must_change_pw=False)
        app.dependency_overrides[get_current_user] = lambda: user
        configured = CloudModeStatus(True, False, "Explicit opt-in required", "(test)", "test-model")
        with (
            patch("backend.api.local_analysis_routes._agentic_check_cloud_mode", side_effect=AssertionError("Cloud must not be contacted")),
            patch("backend.api.local_analysis_routes._agentic_cloud_mode_configuration", return_value=configured),
            TestClient(app) as client,
        ):
            response = client.post(
                "/api/local/ai-business-analyst/agentic",
                data={"question": "Total Net Sales by Department"},
                files={"files": ("sales.csv", b"Department,Net Sales\nRetail,12.5\nRetail,7.5\nOnline,30\n", "text/csv")},
            )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["mode"], "local_fallback")
        self.assertEqual(body["answer"]["mode"], "tabular")
        self.assertIn("Retail: 20.00", body["answer"]["key_evidence"][0]["value"])
        self.assertEqual(body["file_details"][0]["rows"], 3)

    def test_chat_api_keeps_authentication_required(self) -> None:
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from backend.api.local_analysis_routes import create_local_analysis_router

        app = FastAPI()
        app.include_router(create_local_analysis_router())
        with TestClient(app) as client:
            response = client.post("/api/local/ai-business-analyst/agentic", data={"question": "Total Sales"})
        self.assertEqual(response.status_code, 401)

    def test_semicolon_csv_and_decimal_comma_values(self) -> None:
        result = self.engine.ingest(b"Department;Sales\nRetail;12,50\nOnline;30,00\n", "sales.csv")
        self.assertTrue(result.success, result.error)
        answer = analyze_local_question("Total Sales", {"sales.csv": result.df})
        self.assertIn("42.50", answer.detail)

    def test_json_lines_records_are_flattened_and_analyzable(self) -> None:
        result = self.engine.ingest(
            b'{"department":"Retail","metrics":{"sales":12.5}}\n'
            b'{"department":"Online","metrics":{"sales":30.0}}\n',
            "sales.jsonl",
        )
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.row_count, 2)
        self.assertIn("metrics.sales", result.df.columns)
        answer = analyze_local_question("Total sales", {"sales.jsonl": result.df})
        self.assertIn("42.50", answer.detail)

    def test_category_filter_without_only_is_applied(self) -> None:
        frame = pd.DataFrame({"Department": ["Retail", "Online"], "Sales": [10, 20]})
        answer = analyze_local_question("Total Sales for Retail", {"sales.csv": frame})
        self.assertEqual(answer.operation, "total")
        self.assertIn("10.00", answer.detail)

    def test_arabic_and_english_financial_aliases_match_the_same_column(self) -> None:
        frame = pd.DataFrame({
            "Department": ["Operations", "Finance", "Operations"],
            "Annual Payroll": [100.0, 200.0, 50.0],
            "Net Income": [60.0, 40.0, 20.0],
            "Expenses": [10.0, 30.0, 15.0],
        })
        arabic = analyze_local_question("إجمالي الرواتب", {"workforce.csv": frame})
        english = analyze_local_question("Total salary", {"workforce.csv": frame})
        self.assertEqual(arabic.operation, "total")
        self.assertEqual(english.operation, "total")
        self.assertIn("350.00", arabic.detail)
        self.assertIn("350.00", english.detail)

    def test_arabic_net_profit_phrase_prefers_specific_net_income_measure(self) -> None:
        frame = pd.DataFrame({
            "Net Income": [60.0, 40.0],
            "Profit": [5.0, 7.0],
        })
        answer = analyze_local_question("إجمالي صافي الأرباح", {"finance.csv": frame})
        self.assertEqual(answer.operation, "total")
        self.assertIn("Net Income", answer.detail)
        self.assertIn("100.00", answer.detail)

    def test_follow_up_grouping_and_category_filter_reuse_only_prior_local_intent(self) -> None:
        frame = pd.DataFrame({
            "Department": ["Operations", "Finance", "Operations"],
            "Payroll": [10.0, 20.0, 30.0],
        })
        grouped_question = resolve_local_follow_up("And by department?", "Total Payroll")
        grouped = analyze_local_question(grouped_question, {"costs.csv": frame})
        self.assertEqual(grouped.operation, "grouped_total")
        self.assertIn("Operations: 40.00", grouped.detail)

        filtered_question = resolve_local_follow_up("What about Finance?", "Total Payroll by department")
        filtered = analyze_local_question(filtered_question, {"costs.csv": frame})
        self.assertEqual(filtered.operation, "grouped_total")
        self.assertEqual(filtered.detail, "Finance: 20.00")
        self.assertEqual(resolve_local_follow_up("What is attrition?", "Total Payroll"), "What is attrition?")

    def test_arabic_follow_up_keeps_local_query_bounded(self) -> None:
        frame = pd.DataFrame({
            "Department": ["التشغيل", "المالية"],
            "Payroll": [100.0, 200.0],
        })
        resolved = resolve_local_follow_up("حسب القسم", "إجمالي الرواتب")
        answer = analyze_local_question(resolved, {"الرواتب.csv": frame})
        self.assertEqual(answer.operation, "grouped_total")
        self.assertIn("200.00", answer.detail)

    def test_multi_intent_returns_separate_auditable_results_and_fails_closed(self) -> None:
        frame = pd.DataFrame({
            "Department": ["Operations", "Finance", "Operations"],
            "Payroll": [10.0, 20.0, 30.0],
        })
        answer = analyze_local_question(
            "Total Payroll by Department; average Payroll by Department",
            {"costs.csv": frame},
        )
        self.assertEqual(answer.operation, "multi_intent")
        self.assertIn("Operations: 40.00", answer.detail)
        self.assertIn("Operations: 20.00", answer.detail)
        self.assertEqual(len(answer.query_audit["subqueries"]), 2)
        self.assertIn("independently", answer.warnings[0])

        arabic = analyze_local_question(
            "إجمالي الرواتب حسب القسم؛ متوسط الرواتب حسب القسم",
            {"costs.csv": frame},
        )
        self.assertEqual(arabic.operation, "multi_intent")
        self.assertIn("Finance: 20.00", arabic.detail)

        partial = analyze_local_question(
            "Total Payroll by Department; average Unknown Metric by Department",
            {"costs.csv": frame},
        )
        self.assertEqual(partial.operation, "clarify")
        self.assertNotIn("Operations: 40.00", partial.detail)
        self.assertIn("No partial", partial.warnings[0])

    def test_schema_typo_suggests_but_never_applies_fuzzy_measure_mapping(self) -> None:
        frame = pd.DataFrame({
            "Annual Payroll": [100.0, 200.0],
            "Headcount": [4, 5],
        })
        answer = analyze_local_question("Total Annul Payrol", {"workforce.csv": frame})
        self.assertEqual(answer.operation, "clarify")
        self.assertTrue(answer.priority)
        self.assertIn("Annual Payroll", answer.detail)
        self.assertIn("suggestion only", answer.detail)
        self.assertIn("No fuzzy schema mapping was applied", answer.warnings[0])

    def test_arabic_top_costs_groups_by_department(self) -> None:
        frame = pd.DataFrame({
            "Department": ["Operations", "Finance", "Operations", "Finance"],
            "Expenses": [20, 30, 80, 10],
        })
        answer = analyze_local_question("أعلى المصاريف حسب الأقسام", {"costs.csv": frame})
        self.assertEqual(answer.operation, "grouped_total")
        self.assertTrue(answer.detail.startswith("Operations: 100.00"))

    def test_generic_amount_fallback_is_explicit_and_unmatched_filter_fails_closed(self) -> None:
        frame = pd.DataFrame({"Department": ["Operations", "Finance"], "Amount": [20, 30]})
        answer = analyze_local_question("أعلى المصاريف", {"costs.csv": frame})
        self.assertEqual(answer.operation, "maximum")
        self.assertIn("30.00", answer.detail)
        self.assertTrue(answer.warnings)

        unmatched = analyze_local_question("Total expenses for Retail", {"costs.csv": frame})
        self.assertEqual(unmatched.operation, "clarify")
        self.assertTrue(unmatched.priority)

    def test_arabic_expired_contract_count_uses_observed_status_value(self) -> None:
        frame = pd.DataFrame({
            "حالة العقد": ["منتهي", "نشط", "منتهي"],
            "القسم": ["المبيعات", "المبيعات", "العمليات"],
        })
        answer = answer_question(
            "كم عقد منتهي؟", {}, use_llm=False,
            uploaded_frames={"contracts.csv": frame},
        )
        self.assertEqual(answer.matched_query_id, "local.contract_expired_count")
        self.assertIn("العقود المنتهية", answer.answer_text)
        self.assertIn("2", answer.answer_text)

    def test_arabic_schema_based_count_without_contract_template(self) -> None:
        frame = pd.DataFrame({"الحالة": ["منتهي", "نشط", "منتهي", None]})
        answer = answer_question(
            "كم سجل حالته منتهي؟", {}, use_llm=False,
            uploaded_frames={"records.csv": frame},
        )
        self.assertEqual(answer.matched_query_id, "local.count")
        self.assertIn("عدد السجلات", answer.answer_text)
        self.assertEqual(answer.query_audit["cohorts"][0]["value"], 2)

    def test_arabic_how_much_measure_question_sums_instead_of_counting_rows(self) -> None:
        frame = pd.DataFrame({"المبيعات": [12.5, 7.5, 30.0]})
        answer = analyze_local_question("كم قيمة المبيعات؟", {"sales.csv": frame})
        self.assertEqual(answer.operation, "total")
        self.assertIn("50.00", answer.detail)

    def test_natural_english_count_filter_is_applied_and_audited(self) -> None:
        frame = pd.DataFrame({
            "Status": ["Expired", "Active", "Expired", "Active"],
            "Department": ["Sales", "Sales", "Operations", "Operations"],
        })
        answer = answer_question(
            "How many records have Expired status?", {}, use_llm=False,
            uploaded_frames={"contracts.csv": frame},
        )
        self.assertEqual(answer.matched_query_id, "local.count")
        self.assertEqual(answer.query_audit["cohorts"][0]["value"], 2)
        self.assertEqual(answer.query_audit["cohorts"][0]["filters"][0]["value"], "Expired")

    def test_grouping_does_not_treat_measure_name_as_category_filter(self) -> None:
        frame = pd.DataFrame({
            "Department": ["Sales", "Retail", "Sales"],
            "Sales": [10, 20, 30],
        })
        answer = analyze_local_question("Total Sales by Department", {"sales.csv": frame})
        self.assertEqual(answer.operation, "grouped_total")
        self.assertIn("Sales: 40.00", answer.detail)
        self.assertIn("Retail: 20.00", answer.detail)

    def test_duplicate_category_values_across_columns_require_a_column_name(self) -> None:
        frame = pd.DataFrame({
            "Department": ["Retail", "Online"],
            "Region": ["Retail", "West"],
            "Sales": [10, 20],
        })
        answer = analyze_local_question("Total Sales for Retail", {"sales.csv": frame})
        self.assertEqual(answer.operation, "clarify")
        self.assertTrue(answer.priority)
        self.assertIn("Possible values", answer.detail)

    def test_unmatched_uploaded_data_question_returns_schema_guidance(self) -> None:
        frame = pd.DataFrame({
            "Department": ["Sales", "Operations"],
            "Net Revenue": [100.0, 80.0],
            "Period": ["2026-01", "2026-02"],
        })
        answer = answer_question(
            "Which teams improved the most?", {}, use_llm=False,
            uploaded_frames={"monthly.csv": frame},
        )
        self.assertEqual(answer.mode, "guidance")
        self.assertEqual(answer.matched_query_id, "local.schema_guide")
        self.assertIn("Net Revenue", answer.key_evidence[0].excerpt)
        self.assertIn("no metric was calculated", " ".join(answer.limitations).lower())

    def test_ambiguous_day_month_dates_require_iso_format(self) -> None:
        frame = pd.DataFrame({"Date": ["01/02/2026", "02/03/2026"], "Sales": [10, 20]})
        answer = analyze_local_question("Sales trend over time", {"sales.csv": frame})
        self.assertEqual(answer.operation, "clarify")
        self.assertIn("YYYY-MM-DD", answer.text)

    def test_empty_image_only_pdf_has_an_explicit_extraction_error(self) -> None:
        result = self.engine.ingest(_simple_searchable_pdf(""), "scan.pdf")
        self.assertFalse(result.success)
        self.assertIn("OCR", result.error)


if __name__ == "__main__":
    unittest.main()
