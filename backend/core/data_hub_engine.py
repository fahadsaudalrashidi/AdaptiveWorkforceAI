"""Request-scoped data profiling, anomaly checks, forecasting, and safe exports.

Uploaded files are parsed in memory and are never persisted by this module.
Responses contain schema and aggregate statistics only; source records are
never returned or included in generated workbooks/PDFs.
"""
from __future__ import annotations

import io
import re
from dataclasses import asdict
from pathlib import Path
from datetime import timedelta
from typing import Any, Iterable

import numpy as np
import pandas as pd

from backend.core.excel_ingestion_engine import ExcelIngestionEngine, IngestionResult
from backend.core.secure_file_ingestion import SecureFileIngestion

SUPPORTED_DATA_HUB_EXTENSIONS = {".csv", ".xlsx", ".xls", ".json", ".jsonl"}
CLEANING_OPTIONS = {
    "trim_text",
    "remove_blank_rows",
    "remove_blank_columns",
    "remove_exact_duplicates",
}
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_REQUEST_BYTES = 50 * 1024 * 1024
MAX_FILES_PER_REQUEST = 10

_ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩٫٬", "0123456789.,")
_DATE_NAME = re.compile(r"(date|time|month|period|timestamp|تاريخ|شهر|فترة)", re.I)
_EXPIRY_NAME = re.compile(
    r"(expir|expiry|expires|valid.?until|contract.?end|end.?date|"
    r"iqama|visa|permit|تاريخ.?الانتهاء|انتهاء|صلاحية|نهاية.?العقد)",
    re.I,
)
_PAYROLL_NAME = re.compile(r"(payroll|salary|salar|gross|net.?pay|wage|overtime|بدل|راتب|أجر)", re.I)
_BUDGET_NAME = re.compile(r"(budget|actual|planned|forecast|ميزانية|فعلي|مخطط)", re.I)
_COST_NAME = re.compile(
    r"(payroll|salary|salar|gross.?pay|net.?pay|wage|overtime|labor|labour|"
    r"cost|expense|spend|expenditure|budget|amount|مصاريف|مصروفات|تكلفة|"
    r"تكاليف|راتب|أجور|أجر|ميزانية)",
    re.I,
)
_SAFE_LABEL_LIMIT = 160


def _finite(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return round(number, 4) if np.isfinite(number) else None


def _clean_numeric(series: pd.Series) -> pd.Series:
    """Parse common financial formatting without changing source cells."""
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")

    def parse(value: Any) -> float:
        if pd.isna(value):
            return np.nan
        text = str(value).translate(_ARABIC_DIGITS).strip()
        negative = text.startswith("(") and text.endswith(")")
        if negative:
            text = text[1:-1].strip()
        text = re.sub(r"(?i)\b(?:SAR|USD|EUR|GBP|AED)\b", "", text)
        for currency in ("ر.س", "ر س", "ريالات", "ريال", "دولار"):
            text = text.replace(currency, "")
        text = re.sub(r"[$£€٪%\s\u00a0]", "", text)
        # Accept comma decimals only when they are unambiguously not groups of
        # three digits; otherwise commas are treated as thousands separators.
        if "," in text and "." not in text:
            if text.count(",") == 1 and len(text.rsplit(",", 1)[1]) in {1, 2}:
                text = text.replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", "")
        if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
            return np.nan
        try:
            number = float(text)
            return -abs(number) if negative else number
        except ValueError:
            return np.nan

    return pd.Series((parse(value) for value in series), index=series.index, dtype="float64")


def _parse_dates(series: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(series):
        return pd.to_datetime(series, errors="coerce")
    try:
        return pd.to_datetime(series, errors="coerce", format="mixed", dayfirst=True)
    except (TypeError, ValueError):
        return pd.to_datetime(series, errors="coerce", dayfirst=True)


def _column_kind(series: pd.Series, name: str) -> str:
    non_null = series.dropna()
    if non_null.empty:
        return "empty"
    if pd.api.types.is_bool_dtype(series):
        return "boolean"
    if pd.api.types.is_numeric_dtype(series):
        return "numeric"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "date"
    sample_text = non_null.head(100).astype(str).str.strip()
    date_shaped = sample_text.str.contains(
        r"(?:\d{4}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{2,4}|"
        r"\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4})",
        case=False,
        regex=True,
    )
    if _DATE_NAME.search(name) and date_shaped.mean() >= 0.8:
        parsed = _parse_dates(sample_text)
        if parsed.notna().mean() >= 0.8:
            return "date"
    numeric = _clean_numeric(non_null)
    if numeric.notna().mean() >= 0.9 and numeric.notna().sum() >= 2:
        return "numeric"
    if non_null.nunique(dropna=True) <= min(20, max(2, int(len(non_null) * 0.1))):
        return "categorical"
    sample = non_null.head(100).astype(str)
    return "text" if sample.str.len().mean() > 80 else "categorical"


def _sensitivity(name: str, scanner: SecureFileIngestion) -> tuple[bool, str | None]:
    classification = scanner._classify_column(str(name))
    if classification.classification in {"pii", "leakage"}:
        return True, classification.classification
    return False, None


def apply_cleaning(
    frame: pd.DataFrame,
    options: Iterable[str] = (),
) -> tuple[pd.DataFrame, list[dict[str, Any]], dict[str, int]]:
    """Apply only requested reversible/in-memory operations and report counts."""
    selected = set(options)
    unknown = selected - CLEANING_OPTIONS
    if unknown:
        raise ValueError(f"Unsupported cleaning option(s): {', '.join(sorted(unknown))}")

    result = frame.copy()
    applied: list[dict[str, Any]] = []
    before_rows, before_columns = len(result), len(result.columns)

    if "trim_text" in selected:
        changed_cells = 0
        for column in result.select_dtypes(include=["object", "string"]).columns:
            original = result[column]
            cleaned = original.map(lambda value: value.strip() if isinstance(value, str) else value)
            changed_cells += int((original.fillna("\0") != cleaned.fillna("\0")).sum())
            result[column] = cleaned
        applied.append({"id": "trim_text", "changed": changed_cells})

    if "remove_blank_rows" in selected:
        blank = result.apply(
            lambda col: col.isna() | col.astype("string").str.strip().eq("").fillna(False)
        ).all(axis=1)
        removed = int(blank.sum())
        result = result.loc[~blank].copy()
        applied.append({"id": "remove_blank_rows", "changed": removed})

    if "remove_blank_columns" in selected:
        blank_cols = [
            col for col in result.columns
            if result[col].isna().all()
            or result[col].astype("string").str.strip().eq("").fillna(False).all()
        ]
        result = result.drop(columns=blank_cols)
        applied.append({"id": "remove_blank_columns", "changed": len(blank_cols)})

    if "remove_exact_duplicates" in selected:
        before = len(result)
        result = result.drop_duplicates(keep="first").reset_index(drop=True)
        applied.append({"id": "remove_exact_duplicates", "changed": before - len(result)})

    counts = {
        "source_rows": before_rows,
        "analyzed_rows": len(result),
        "source_columns": before_columns,
        "analyzed_columns": len(result.columns),
    }
    return result.reset_index(drop=True), applied, counts


def _outlier_count(values: pd.Series) -> tuple[int, float | None, float | None]:
    values = values.dropna()
    if len(values) < 8:
        return 0, None, None
    q1, q3 = values.quantile([0.25, 0.75])
    iqr = float(q3 - q1)
    if not np.isfinite(iqr) or iqr <= 0:
        return 0, None, None
    low, high = float(q1 - 1.5 * iqr), float(q3 + 1.5 * iqr)
    return int(((values < low) | (values > high)).sum()), _finite(low), _finite(high)


def _payroll_checks(frame: pd.DataFrame, sensitive: set[str]) -> list[dict[str, Any]]:
    """Return only cohort-level consistency counts; never expose row values."""
    if len(frame) < 10:
        return []
    columns = {str(c).strip().lower(): c for c in frame.columns}
    gross = next((c for key, c in columns.items() if "gross" in key and _PAYROLL_NAME.search(key)), None)
    net = next((c for key, c in columns.items() if "net" in key and _PAYROLL_NAME.search(key)), None)
    checks: list[dict[str, Any]] = []

    if gross is not None and net is not None:
        gross_values, net_values = _clean_numeric(frame[gross]), _clean_numeric(frame[net])
        valid = gross_values.notna() & net_values.notna()
        evaluated = int(valid.sum())
        if evaluated >= 10:
            violations = int((gross_values[valid] < net_values[valid]).sum())
            if violations:
                checks.append({
                    "id": "gross_below_net",
                    "kind": "payroll_consistency",
                    "column": f"{gross} / {net}",
                    "count": violations,
                    "evaluated": evaluated,
                    "percentage": _finite(violations * 100 / evaluated),
                    "severity": "high",
                    "method": "Gross pay is below net pay",
                })
    return checks


def _compliance_checks(frame: pd.DataFrame) -> list[dict[str, Any]]:
    checks: list[dict[str, Any]] = []
    today = pd.Timestamp(pd.Timestamp.now(tz="Asia/Riyadh").date())
    for column in frame.columns:
        if not _EXPIRY_NAME.search(str(column)):
            continue
        dates = _parse_dates(frame[column]).dropna()
        if dates.empty:
            continue
        if isinstance(dates.dtype, pd.DatetimeTZDtype):
            dates = dates.dt.tz_localize(None)
        expired = int((dates.dt.normalize() < today).sum())
        due_soon = int(((dates.dt.normalize() >= today) & (dates.dt.normalize() <= today + timedelta(days=30))).sum())
        if expired or due_soon:
            checks.append({
                "id": "compliance_expiry",
                "kind": "compliance",
                "column": str(column),
                "count": expired + due_soon,
                "expired": expired,
                "due_30_days": due_soon,
                "evaluated": int(len(dates)),
                "severity": "high" if expired else "moderate",
                "method": "Expiry dates compared with the server's current date",
            })
    return checks


def _temporal_cost_checks(frame: pd.DataFrame, sensitive: set[str]) -> list[dict[str, Any]]:
    """Flag sharp rises in monthly aggregate costs using a rolling robust baseline."""
    date_columns = [
        column for column in frame.columns
        if _DATE_NAME.search(str(column))
        or pd.api.types.is_datetime64_any_dtype(frame[column])
    ]
    if not date_columns:
        return []
    date_column = date_columns[0]
    dates = _parse_dates(frame[date_column])
    if isinstance(dates.dtype, pd.DatetimeTZDtype):
        dates = dates.dt.tz_localize(None)

    checks: list[dict[str, Any]] = []
    for column in frame.columns:
        name = str(column)
        if name in sensitive or not _COST_NAME.search(name):
            continue
        values = _clean_numeric(frame[column])
        valid = dates.notna() & values.notna()
        if int(valid.sum()) < 8:
            continue
        data = pd.DataFrame({"period": dates[valid].dt.to_period("M"), "value": values[valid]})
        monthly = data.groupby("period", sort=True)["value"].sum(min_count=1).dropna()
        if len(monthly) < 8:
            continue

        flagged_periods: list[str] = []
        flagged_indices: list[int] = []
        window = 6
        for index in range(window, len(monthly)):
            history = monthly.iloc[index - window:index]
            baseline = float(history.median())
            mad = float((history - baseline).abs().median())
            robust_scale = 1.4826 * mad
            # A 50% floor avoids treating minor jitter as an incident when
            # historical values are flat and MAD is zero.
            threshold = baseline + max(3 * robust_scale, abs(baseline) * 0.5, 1e-9)
            if float(monthly.iloc[index]) > threshold:
                flagged_periods.append(str(monthly.index[index]))
                flagged_indices.append(index)
        evaluated = len(monthly) - window
        if not flagged_periods:
            continue
        # Preserve the existing detector's flags and score exactly. This
        # sidecar only adds context when the same month was independently
        # flagged in each of the prior two annual cycles.
        flagged_index_set = set(flagged_indices)
        recurring_periods: list[str] = []
        for index in flagged_indices:
            if index < 24:
                continue
            period, prior_year, two_years_prior = (
                monthly.index[index], monthly.index[index - 12], monthly.index[index - 24]
            )
            if (
                period.month == prior_year.month == two_years_prior.month
                and period.year - prior_year.year == 1
                and prior_year.year - two_years_prior.year == 1
                and index - 12 in flagged_index_set
                and index - 24 in flagged_index_set
            ):
                recurring_periods.append(str(period))
        rate = len(flagged_periods) * 100 / max(evaluated, 1)
        checks.append({
            "id": "temporal_cost_spike",
            "kind": "temporal_cost_spike",
            "column": name[:_SAFE_LABEL_LIMIT],
            "count": len(flagged_periods),
            "evaluated": evaluated,
            "percentage": _finite(rate),
            "severity": "high" if rate >= 20 else "moderate",
            "method": "Monthly aggregate compared with the prior six-month median + robust MAD threshold (50% minimum rise)",
            "method_ar": "مقارنة الإجمالي الشهري بوسيط الأشهر الستة السابقة وحدّ MAD متين (ارتفاع أدنى 50%)",
            "periods": flagged_periods[:12],
            "seasonal_pattern_periods": recurring_periods[:12],
            "seasonality_context": (
                "Same calendar month was also flagged in both prior annual cycles; check the payroll calendar before treating it as a new incident."
                if recurring_periods else
                "Seasonal recurrence is not confirmed; treat this as a screening signal, not a conclusion."
            ),
            "seasonality_context_ar": (
                "رُصدت الإشارة أيضاً في الشهر نفسه خلال الدورتين السنويتين السابقتين؛ راجع تقويم الرواتب قبل اعتبارها حالة جديدة."
                if recurring_periods else
                "لم يتأكد التكرار الموسمي؛ تعامل مع هذه الإشارة كتنبيه للفحص لا كاستنتاج."
            ),
        })
    return checks


def profile_frame(
    frame: pd.DataFrame,
    *,
    filename: str,
    sheet: str,
    cleaning_options: Iterable[str] = (),
) -> dict[str, Any]:
    scanner = SecureFileIngestion()
    cleaned, applied, counts = apply_cleaning(frame, cleaning_options)
    numeric_columns: list[str] = []
    date_columns: list[str] = []
    columns: list[dict[str, Any]] = []
    anomalies: list[dict[str, Any]] = []
    sensitive_columns: set[str] = set()
    missing_cells = 0

    for raw_name in cleaned.columns:
        name = str(raw_name)
        series = cleaned[raw_name]
        missing = int(series.isna().sum())
        missing_cells += missing
        kind = _column_kind(series, name)
        sensitive, sensitive_reason = _sensitivity(name, scanner)
        if sensitive:
            sensitive_columns.add(name)
        entry: dict[str, Any] = {
            "name": name[:_SAFE_LABEL_LIMIT],
            "dtype": str(series.dtype)[:40],
            "type": kind,
            "missing_count": missing,
            "missing_percent": _finite(missing * 100 / max(len(cleaned), 1)),
            "unique_count": int(series.nunique(dropna=True)),
            "sensitive": sensitive,
            "sensitive_reason": sensitive_reason,
        }
        if kind == "date":
            date_columns.append(name)
        if kind == "numeric":
            numeric = _clean_numeric(series)
            if not sensitive:
                valid = numeric.dropna()
                outliers, low, high = _outlier_count(valid)
                entry["numeric_summary"] = {
                    "non_null": int(len(valid)),
                    "mean": _finite(valid.mean()) if len(valid) else None,
                    "median": _finite(valid.median()) if len(valid) else None,
                    "min": _finite(valid.min()) if len(valid) else None,
                    "max": _finite(valid.max()) if len(valid) else None,
                    "std": _finite(valid.std()) if len(valid) > 1 else None,
                }
                entry["outlier_count"] = outliers
                numeric_columns.append(name)
                if outliers:
                    rate = outliers * 100 / max(len(valid), 1)
                    anomalies.append({
                        "id": "numeric_outliers",
                        "kind": "distribution",
                        "column": name[:_SAFE_LABEL_LIMIT],
                        "count": outliers,
                        "evaluated": int(len(valid)),
                        "percentage": _finite(rate),
                        "severity": "high" if rate >= 10 else "moderate",
                        "method": "Tukey IQR fence (1.5 × IQR)",
                        "lower_fence": low,
                        "upper_fence": high,
                    })
            else:
                entry["numeric_summary"] = None
                entry["outlier_count"] = None
        columns.append(entry)

    anomalies.extend(_payroll_checks(cleaned, sensitive_columns))
    anomalies.extend(_compliance_checks(cleaned))
    anomalies.extend(_temporal_cost_checks(cleaned, sensitive_columns))

    duplicate_rows = int(cleaned.duplicated().sum())
    suggestions = [
        {"id": "trim_text", "eligible": True, "reason": "Trim leading/trailing spaces in text cells."},
        {"id": "remove_blank_rows", "eligible": int(cleaned.isna().all(axis=1).sum()), "reason": "Remove rows with no values."},
        {
            "id": "remove_blank_columns",
            "eligible": int(sum(
                col.isna().all() or col.astype("string").str.strip().eq("").fillna(False).all()
                for _, col in cleaned.items()
            )),
            "reason": "Remove columns containing no values.",
        },
        {"id": "remove_exact_duplicates", "eligible": duplicate_rows, "reason": "Remove exact duplicate rows only."},
    ]
    denominator = max(len(cleaned) * max(len(cleaned.columns), 1), 1)
    missing_percent = missing_cells * 100 / denominator
    duplicate_percent = duplicate_rows * 100 / max(len(cleaned), 1)
    quality = max(0.0, min(100.0, 100 - min(50, missing_percent * 0.75) - min(35, duplicate_percent * 0.5)))

    return {
        "id": f"{Path(filename).name}:{sheet}"[:240],
        "filename": Path(filename).name[:_SAFE_LABEL_LIMIT],
        "sheet": str(sheet)[:_SAFE_LABEL_LIMIT],
        "source_type": Path(filename).suffix.lower().lstrip("."),
        **counts,
        "column_count": len(columns),
        "missing_cells": missing_cells,
        "missing_percent": _finite(missing_percent),
        "duplicate_rows": duplicate_rows,
        "quality_score": _finite(quality),
        "columns": columns,
        "numeric_columns": numeric_columns,
        "date_columns": date_columns,
        "anomalies": anomalies,
        "cleaning_suggestions": suggestions,
        "applied_cleaning": applied,
        "warnings": [
            "Column names identified as sensitive are shown without value-level statistics.",
            "No source rows or individual employee values are returned or stored.",
            "Anomaly flags are review signals, not findings of misconduct or errors.",
        ],
    }


def profile_upload(
    raw: bytes,
    filename: str,
    cleaning_options: Iterable[str] = (),
) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_DATA_HUB_EXTENSIONS:
        raise ValueError(f"Unsupported file type {suffix or '(no extension)'}. Use CSV, XLS, XLSX, or JSON.")
    if not raw:
        raise ValueError("The uploaded file is empty.")
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError("Each file must be 25 MB or smaller.")
    engine = ExcelIngestionEngine()
    results = engine.ingest_all(raw, Path(filename).name)
    profiles: list[dict[str, Any]] = []
    for sheet, result in results:
        if not result.success or result.df is None:
            raise ValueError(result.error or f"Could not read {sheet}.")
        if result.df.empty:
            continue
        profile = profile_frame(
            result.df, filename=filename, sheet=sheet, cleaning_options=cleaning_options
        )
        profile["warnings"] = list(dict.fromkeys(profile["warnings"] + result.warnings))
        profile["ingestion_transformations"] = result.transformations
        profiles.append(profile)
    if not profiles:
        raise ValueError("No non-empty tables or worksheets were found.")
    return profiles


def forecast_frame(
    frame: pd.DataFrame,
    *,
    filename: str,
    sheet: str,
    date_column: str,
    metric_column: str,
    aggregation: str = "sum",
    horizon: int = 3,
) -> dict[str, Any]:
    if aggregation not in {"sum", "mean", "count"}:
        raise ValueError("Aggregation must be sum, mean, or count.")
    if not 1 <= int(horizon) <= 12:
        raise ValueError("Forecast horizon must be between 1 and 12 months.")
    if date_column not in frame.columns:
        raise ValueError("Choose a date column present in the selected worksheet.")
    if metric_column != "__rows__" and metric_column not in frame.columns:
        raise ValueError("Choose a measure column present in the selected worksheet.")
    if metric_column == "__rows__" and aggregation != "count":
        raise ValueError("Record count requires the count aggregation.")
    if metric_column != "__rows__":
        sensitive, _ = _sensitivity(metric_column, SecureFileIngestion())
        if sensitive:
            raise ValueError("Forecasts are disabled for columns identified as personal or post-event data.")

    dates = _parse_dates(frame[date_column])
    if isinstance(dates.dtype, pd.DatetimeTZDtype):
        dates = dates.dt.tz_localize(None)
    values = pd.Series(1.0, index=frame.index) if metric_column == "__rows__" else _clean_numeric(frame[metric_column])
    valid = dates.notna() & values.notna()
    if not valid.any():
        raise ValueError("No rows have both a valid date and a numeric measure.")
    data = pd.DataFrame({"date": dates[valid], "value": values[valid]})
    data["period"] = data["date"].dt.to_period("M")
    grouped = data.groupby("period")["value"]
    if aggregation == "sum":
        monthly = grouped.sum()
    elif aggregation == "mean":
        monthly = grouped.mean()
    else:
        monthly = grouped.count().astype(float)
    monthly = monthly.sort_index()
    monthly = monthly[monthly.notna()]
    if len(monthly) < 8:
        raise ValueError(
            f"At least 8 monthly observations are required for a trend projection; found {len(monthly)}."
        )

    y = monthly.to_numpy(dtype=float)
    periods = list(monthly.index)

    def predict(model: str, train_periods: list[pd.Period], train_values: np.ndarray, target: pd.Period) -> float:
        if model == "seasonal_naive":
            prior_year = target - 12
            prior = next(
                (value for period, value in zip(train_periods, train_values) if period == prior_year),
                None,
            )
            if prior is not None:
                return float(prior)
        if model == "linear_trend" and len(train_values) >= 2:
            offsets = np.array([period.ordinal - train_periods[0].ordinal for period in train_periods], dtype=float)
            slope, intercept = np.polyfit(offsets, train_values, 1)
            return float(intercept + slope * (target.ordinal - train_periods[0].ordinal))
        return float(train_values[-1])

    # Select among transparent local baselines using one-step rolling-origin
    # validation. Seasonal-naive is offered only when enough history exists.
    min_train = max(4, min(12, len(y) - 2))
    candidate_models = ["naive", "linear_trend"]
    if len(y) >= 18:
        candidate_models.append("seasonal_naive")
    errors_by_model: dict[str, list[float]] = {model: [] for model in candidate_models}
    for origin in range(min_train, len(y)):
        for model in candidate_models:
            estimate = predict(model, periods[:origin], y[:origin], periods[origin])
            errors_by_model[model].append(float(y[origin] - estimate))
    model_order = {model: index for index, model in enumerate(candidate_models)}
    model_mae = {
        model: float(np.mean(np.abs(errors)))
        for model, errors in errors_by_model.items()
        if errors
    }
    selected_model = min(model_mae, key=lambda model: (model_mae[model], model_order[model]))
    validation_errors = np.asarray(errors_by_model[selected_model], dtype=float)
    backtest_rmse = float(np.sqrt(np.mean(validation_errors ** 2))) if len(validation_errors) else 0.0
    # A non-zero scale floor prevents a short, exactly linear sample from
    # presenting false certainty through a zero-width interval.
    scale_floor = max(float(np.median(np.abs(y))) * 0.05, 1e-6)
    error_scale = max(backtest_rmse, scale_floor)

    projection: list[dict[str, Any]] = []
    for offset in range(1, int(horizon) + 1):
        period = monthly.index[-1] + offset
        estimate = predict(selected_model, periods, y, period)
        margin = 1.96 * error_scale * float(np.sqrt(offset))
        if metric_column == "__rows__" or aggregation == "count":
            estimate = max(0.0, estimate)
        projection.append({
            "period": str(period),
            "estimate": _finite(estimate),
            "lower": _finite(max(0.0, estimate - margin) if metric_column == "__rows__" or aggregation == "count" else estimate - margin),
            "upper": _finite(estimate + margin),
        })

    return {
        "filename": Path(filename).name[:_SAFE_LABEL_LIMIT],
        "sheet": str(sheet)[:_SAFE_LABEL_LIMIT],
        "date_column": date_column[:_SAFE_LABEL_LIMIT],
        "metric_column": metric_column if metric_column == "__rows__" else metric_column[:_SAFE_LABEL_LIMIT],
        "aggregation": aggregation,
        "period_grain": "month",
        "method": f"Rolling-origin model selection ({selected_model.replace('_', ' ')}) with widening empirical error bounds",
        "selected_model": selected_model,
        "backtest_mae": _finite(model_mae[selected_model]),
        "validation_periods": len(validation_errors),
        "observation_count": int(len(monthly)),
        "observed": [{"period": str(period), "value": _finite(value)} for period, value in monthly.items()],
        "forecast": projection,
        "limitations": [
            "The model is selected using a short one-step rolling-origin backtest; selection may be unstable with limited history.",
            "Bounds are approximate empirical error ranges with a non-zero scale floor, not guarantees or calibrated probabilities.",
            "The models do not incorporate causal drivers, operational constraints, or future policy changes.",
            "Validate against domain knowledge before using in workforce or financial decisions.",
        ],
    }


def workforce_scenario_frame(
    frame: pd.DataFrame,
    *,
    filename: str,
    sheet: str,
    department_column: str,
    headcount_column: str,
    payroll_column: str,
    overtime_column: str | None,
    scenarios: list[dict[str, Any]],
) -> dict[str, Any]:
    """Run transparent what-if calculations and return aggregates only."""
    if frame.empty:
        raise ValueError("The selected worksheet has no data rows.")
    if not scenarios or len(scenarios) > 3:
        raise ValueError("Provide one to three scenarios.")
    if department_column and department_column not in frame.columns:
        raise ValueError("Choose a department or workforce-group column in the selected worksheet.")
    for column in (payroll_column,):
        if column not in frame.columns:
            raise ValueError("Choose a payroll measure present in the selected worksheet.")
    if headcount_column != "__rows__" and headcount_column not in frame.columns:
        raise ValueError("Choose a headcount measure present in the selected worksheet.")
    if overtime_column and overtime_column not in frame.columns:
        raise ValueError("Choose an overtime measure present in the selected worksheet.")
    if len(set(map(str, frame.columns))) != len(frame.columns):
        raise ValueError("Duplicate column names must be resolved before scenario modeling.")
    measure_columns = [
        column for column in (headcount_column, payroll_column, overtime_column)
        if column and column != "__rows__"
    ]
    if len(set(measure_columns)) != len(measure_columns):
        raise ValueError("Choose separate headcount, payroll, and overtime measures.")

    scanner = SecureFileIngestion()
    selected_columns = [payroll_column]
    if department_column:
        selected_columns.append(department_column)
    if headcount_column != "__rows__":
        selected_columns.append(headcount_column)
    if overtime_column:
        selected_columns.append(overtime_column)
    for column in selected_columns:
        sensitive, _ = _sensitivity(column, scanner)
        if sensitive:
            raise ValueError("Scenario modeling is disabled for columns identified as personal or post-event data.")

    if department_column and department_column in {payroll_column, headcount_column, overtime_column}:
        raise ValueError("Choose separate department, headcount, payroll, and overtime columns.")

    departments = (
        frame[department_column].astype("string").str.strip()
        if department_column
        else pd.Series("All workforce", index=frame.index, dtype="string")
    )
    headcount = (
        pd.Series(1.0, index=frame.index, dtype="float64")
        if headcount_column == "__rows__"
        else _clean_numeric(frame[headcount_column])
    )
    payroll = _clean_numeric(frame[payroll_column])
    overtime = (
        _clean_numeric(frame[overtime_column])
        if overtime_column
        else pd.Series(0.0, index=frame.index, dtype="float64")
    )
    valid = departments.notna() & departments.ne("") & headcount.notna() & payroll.notna() & overtime.notna()
    if not valid.any():
        raise ValueError("No rows contain a valid workforce group, headcount, and payroll measure.")
    if (headcount[valid] < 0).any() or (payroll[valid] < 0).any() or (overtime[valid] < 0).any():
        raise ValueError("Headcount and cost measures must be non-negative for scenario modeling.")

    data = pd.DataFrame({
        "department": departments[valid].astype(str).str.slice(0, _SAFE_LABEL_LIMIT),
        "headcount": headcount[valid].astype(float),
        "payroll": payroll[valid].astype(float),
        "overtime": overtime[valid].astype(float),
    })
    if data["department"].nunique(dropna=True) > 500:
        raise ValueError("The selected workforce group column has too many distinct values to model safely.")
    grouped = data.groupby("department", sort=True, dropna=True).sum(numeric_only=True)
    if grouped.empty or grouped["headcount"].sum() <= 0:
        raise ValueError("A positive headcount total is required for scenario modeling.")

    def totals_for(group_values: pd.DataFrame) -> dict[str, float | None]:
        payroll_total = float(group_values["payroll"].sum())
        overtime_total = float(group_values["overtime"].sum()) if overtime_column else None
        return {
            "headcount": _finite(group_values["headcount"].sum()) or 0.0,
            "payroll": _finite(payroll_total) or 0.0,
            "overtime": _finite(overtime_total) if overtime_total is not None else None,
            "labor_cost": _finite(payroll_total + (overtime_total or 0.0)) or 0.0,
        }

    def public_groups(group_values: pd.DataFrame, withheld: bool) -> list[dict[str, Any]]:
        if withheld or not department_column:
            return []
        result: list[dict[str, Any]] = []
        for label, row in group_values.iterrows():
            payroll_value = float(row["payroll"])
            overtime_value = float(row["overtime"]) if overtime_column else None
            result.append({
                "label": str(label)[:_SAFE_LABEL_LIMIT],
                "headcount": _finite(row["headcount"]) or 0.0,
                "payroll": _finite(payroll_value) or 0.0,
                "overtime": _finite(overtime_value) if overtime_value is not None else None,
                "labor_cost": _finite(payroll_value + (overtime_value or 0.0)) or 0.0,
            })
        return result

    withheld_groups = bool(department_column and (grouped["headcount"] < 5).any())
    baseline_totals = totals_for(grouped)
    baseline = {
        **baseline_totals,
        "groups": public_groups(grouped, withheld_groups),
        "group_details_withheld": withheld_groups,
    }

    scenarios_result: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for index, scenario in enumerate(scenarios, start=1):
        if not isinstance(scenario, dict):
            raise ValueError("Each scenario must be an object.")
        scenario_id = str(scenario.get("id", f"scenario-{index}"))[:40]
        name = str(scenario.get("name", f"Scenario {index}")).strip()[:60]
        if not scenario_id or scenario_id in seen_ids or not name:
            raise ValueError("Scenario names and IDs must be non-empty and IDs must be unique.")
        seen_ids.add(scenario_id)
        try:
            headcount_change = float(scenario["headcount_change_pct"])
            payroll_growth = float(scenario["payroll_growth_pct"])
            overtime_reduction = float(scenario["overtime_reduction_pct"])
        except (KeyError, TypeError, ValueError, OverflowError) as exc:
            raise ValueError("Each scenario must include valid headcount, payroll, and overtime assumptions.") from exc
        if not all(np.isfinite(value) for value in (headcount_change, payroll_growth, overtime_reduction)):
            raise ValueError("Scenario assumptions must be finite numbers.")
        if not -100 <= headcount_change <= 200:
            raise ValueError("Headcount change must be between -100% and 200%.")
        if not -100 <= payroll_growth <= 200:
            raise ValueError("Payroll growth must be between -100% and 200%.")
        if not -100 <= overtime_reduction <= 100:
            raise ValueError("Overtime reduction must be between -100% and 100%.")

        headcount_factor = 1 + headcount_change / 100
        payroll_factor = 1 + payroll_growth / 100
        overtime_factor = 1 - overtime_reduction / 100
        projected = grouped.copy()
        projected["headcount"] *= headcount_factor
        projected["payroll"] *= headcount_factor * payroll_factor
        projected["overtime"] *= headcount_factor * overtime_factor
        totals = totals_for(projected)
        scenarios_result.append({
            "id": scenario_id,
            "name": name,
            "assumptions": {
                "headcount_change_pct": _finite(headcount_change),
                "payroll_growth_pct": _finite(payroll_growth),
                "overtime_reduction_pct": _finite(overtime_reduction),
            },
            "totals": totals,
            "groups": public_groups(projected, withheld_groups),
            "group_details_withheld": withheld_groups,
        })

    excluded_rows = int(len(frame) - int(valid.sum()))
    limitations = [
        "These are deterministic what-if calculations, not predictive forecasts; the supplied percentages are planning assumptions.",
        "Headcount changes are applied proportionally across groups. Payroll is scaled by the headcount factor and the separate payroll-growth assumption.",
        "Overtime is scaled by projected headcount and the separate overtime-reduction assumption.",
        "This calculation does not infer labor-law or Saudization requirements and is not compliance advice.",
    ]
    if withheld_groups:
        limitations.append(
            "Department-level details are withheld because at least one group has fewer than five people; overall totals still include all valid groups."
        )

    return {
        "filename": Path(filename).name[:_SAFE_LABEL_LIMIT],
        "sheet": str(sheet)[:_SAFE_LABEL_LIMIT],
        "department_column": department_column[:_SAFE_LABEL_LIMIT],
        "headcount_column": headcount_column,
        "payroll_column": payroll_column[:_SAFE_LABEL_LIMIT],
        "overtime_column": overtime_column[:_SAFE_LABEL_LIMIT] if overtime_column else None,
        "source_rows": int(len(frame)),
        "analyzed_rows": int(valid.sum()),
        "excluded_rows": excluded_rows,
        "baseline": baseline,
        "scenarios": scenarios_result,
        "limitations": limitations,
        "privacy": {
            "mode": "request_scoped_local_processing",
            "persisted": False,
            "source_rows_returned": False,
        },
    }


def workforce_skills_frame(
    supply_frame: pd.DataFrame,
    *,
    filename: str,
    sheet: str,
    supply_role_column: str,
    supply_skill_column: str,
    supply_count_column: str | None = None,
    demand_frame: pd.DataFrame | None = None,
    demand_filename: str | None = None,
    demand_sheet: str | None = None,
    demand_role_column: str | None = None,
    demand_skill_column: str | None = None,
    demand_count_column: str | None = None,
) -> dict[str, Any]:
    """Compare role/skill supply with explicit demand, returning safe aggregates only."""
    if supply_frame.empty:
        raise ValueError("The selected skills inventory has no data rows.")

    def validate_columns(
        frame: pd.DataFrame,
        columns: list[str],
        context: str,
    ) -> None:
        if len(set(map(str, frame.columns))) != len(frame.columns):
            raise ValueError(f"Duplicate column names in the {context} must be resolved before skills analysis.")
        missing = [column for column in columns if column not in frame.columns]
        if missing:
            raise ValueError(f"Choose mapped role, skill, and count columns present in the {context}.")
        scanner = SecureFileIngestion()
        if any(_sensitivity(column, scanner)[0] for column in columns):
            raise ValueError(f"Skills analysis is disabled for {context} columns identified as personal or post-event data.")

    supply_columns = [supply_role_column, supply_skill_column]
    if supply_count_column:
        supply_columns.append(supply_count_column)
    validate_columns(supply_frame, supply_columns, "skills inventory")
    if supply_role_column == supply_skill_column:
        raise ValueError("Choose separate role and skill columns in the skills inventory.")
    if supply_count_column in {supply_role_column, supply_skill_column}:
        raise ValueError("Choose a separate numeric count column in the skills inventory.")

    demand_present = demand_frame is not None
    if demand_present:
        if demand_frame.empty:
            raise ValueError("The selected role/skill demand table has no data rows.")
        if not demand_role_column or not demand_skill_column or not demand_count_column:
            raise ValueError("Map role, skill, and numeric demand columns explicitly.")
        demand_columns = [demand_role_column, demand_skill_column, demand_count_column]
        validate_columns(demand_frame, demand_columns, "role/skill demand table")
        if len(set(demand_columns)) != len(demand_columns):
            raise ValueError("Choose separate role, skill, and demand columns.")

    def normalize_labels(frame: pd.DataFrame, role_col: str, skill_col: str) -> pd.DataFrame:
        roles = frame[role_col].astype("string").str.strip()
        skills = frame[skill_col].astype("string").str.strip()
        valid = roles.notna() & roles.ne("") & skills.notna() & skills.ne("")
        result = pd.DataFrame({
            "role": roles[valid].astype(str).str.slice(0, _SAFE_LABEL_LIMIT),
            "skill": skills[valid].astype(str).str.slice(0, _SAFE_LABEL_LIMIT),
        })
        result["role_key"] = result["role"].str.casefold()
        result["skill_key"] = result["skill"].str.casefold()
        return result

    supply = normalize_labels(supply_frame, supply_role_column, supply_skill_column)
    if supply.empty:
        raise ValueError("No inventory rows contain both a role and a skill.")
    supply["supply_records"] = 1.0
    supply["supply"] = (
        _clean_numeric(supply_frame.loc[supply.index, supply_count_column]).astype(float)
        if supply_count_column
        else 1.0
    )
    if supply["supply"].isna().any():
        raise ValueError("The mapped supply count must contain a numeric value for every role/skill row.")
    if (supply["supply"] < 0).any():
        raise ValueError("Supply counts must be non-negative.")
    supply_grouped = supply.groupby(["role_key", "skill_key"], sort=True).agg(
        role=("role", "first"),
        skill=("skill", "first"),
        supply=("supply", "sum"),
        supply_records=("supply_records", "sum"),
    )
    if len(supply_grouped) > 500:
        raise ValueError("The selected skills inventory has too many role/skill groups to analyze safely.")

    demand_grouped = pd.DataFrame(
        columns=["demand_role", "demand_skill", "demand", "demand_records"],
        index=pd.MultiIndex.from_arrays([[], []], names=["role_key", "skill_key"]),
    )
    if demand_present and demand_frame is not None and demand_role_column and demand_skill_column and demand_count_column:
        demand = normalize_labels(demand_frame, demand_role_column, demand_skill_column)
        demand["demand"] = _clean_numeric(demand_frame.loc[demand.index, demand_count_column]).astype(float)
        if demand["demand"].isna().any():
            raise ValueError("The mapped demand count must contain a numeric value for every role/skill row.")
        if (demand["demand"] < 0).any():
            raise ValueError("Demand counts must be non-negative.")
        demand_grouped = demand.groupby(["role_key", "skill_key"], sort=True).agg(
            demand_role=("role", "first"),
            demand_skill=("skill", "first"),
            demand=("demand", "sum"),
            demand_records=("demand", "size"),
        )
        if len(demand_grouped) > 500:
            raise ValueError("The selected demand table has too many role/skill groups to analyze safely.")

    combined = supply_grouped.join(demand_grouped, how="outer")
    combined["role"] = combined["role"].fillna(combined["demand_role"])
    combined["skill"] = combined["skill"].fillna(combined["demand_skill"])
    combined["supply"] = combined["supply"].fillna(0.0)
    combined["supply_records"] = combined["supply_records"].fillna(0.0)
    if demand_present:
        combined["demand"] = combined["demand"].fillna(0.0)
        combined["demand_records"] = combined["demand_records"].fillna(0.0)
        combined["privacy_cohort"] = combined[["supply_records", "demand"]].min(axis=1)
        combined["gap"] = (combined["demand"] - combined["supply"]).clip(lower=0)
        combined["surplus"] = (combined["supply"] - combined["demand"]).clip(lower=0)
    else:
        combined["privacy_cohort"] = combined["supply_records"]
        combined["demand"] = np.nan
        combined["demand_records"] = 0.0
        combined["gap"] = np.nan
        combined["surplus"] = np.nan

    if combined.empty or combined["supply"].sum() <= 0:
        raise ValueError("A positive skills supply total is required for skills analysis.")
    suppressed = bool((combined["privacy_cohort"] < 5).any())
    enough_overall = float(combined["supply_records"].sum()) >= 5
    enough_demand_overall = not demand_present or float(combined["demand"].sum()) >= 5
    totals_suppressed = not enough_overall or not enough_demand_overall
    can_show_totals = not totals_suppressed and not suppressed

    skill_groups: list[dict[str, Any]] = []
    if demand_present:
        per_skill = combined.groupby("skill_key", sort=True)
        for _, rows in per_skill:
            if (rows["privacy_cohort"] < 5).any():
                suppressed = True
                continue
            demand_total = float(rows["demand"].sum())
            supply_total = float(rows["supply"].sum())
            surplus_total = float(rows["surplus"].sum())
            gap_total = float(rows["gap"].sum())
            skill_groups.append({
                "skill": str(rows["skill"].iloc[0])[:_SAFE_LABEL_LIMIT],
                "supply": _finite(supply_total) or 0.0,
                "demand": _finite(demand_total) or 0.0,
                "coverage_pct": _finite(supply_total / demand_total * 100) if demand_total > 0 else None,
                "gap": _finite(gap_total) or 0.0,
                "redeployable_units": _finite(min(surplus_total, gap_total)) or 0.0,
            })

    detail_groups: list[dict[str, Any]] = []
    if not suppressed and not totals_suppressed:
        for _, row in combined.iterrows():
            demand_value = _finite(row["demand"]) if demand_present else None
            supply_value = _finite(row["supply"]) or 0.0
            detail_groups.append({
                "role": str(row["role"])[:_SAFE_LABEL_LIMIT],
                "skill": str(row["skill"])[:_SAFE_LABEL_LIMIT],
                "supply": supply_value,
                "demand": demand_value,
                "coverage_pct": (
                    _finite(supply_value / demand_value * 100)
                    if demand_value is not None and demand_value > 0
                    else None
                ),
                "gap": _finite(row["gap"]) if demand_present else None,
                "surplus": _finite(row["surplus"]) if demand_present else None,
            })
    else:
        skill_groups = []

    total_supply = float(combined["supply"].sum())
    total_demand = float(combined["demand"].sum()) if demand_present else None
    total_gap = float(combined["gap"].sum()) if demand_present else None
    total_surplus = float(combined["surplus"].sum()) if demand_present else None
    redeployable = (
        float(sum(group["redeployable_units"] for group in skill_groups))
        if demand_present and can_show_totals and not suppressed
        else None
    )
    limitations = [
        "Supply counts are summed from the explicitly mapped column, or counted as one per inventory row when row counting is selected.",
        "Coverage is supply divided by explicitly supplied demand; no demand is inferred from inventory.",
        "Possible redeployment is an aggregate upper bound: per-skill role surpluses are matched to role gaps. It does not recommend people or individual moves.",
        "Role/skill records are treated as one skill per row; labels are matched case-insensitively after trimming.",
    ]
    if suppressed or totals_suppressed:
        limitations.append(
            "Role/skill breakdowns or totals were withheld because at least one evaluated cohort is below five; no person-level records are returned."
        )
    if not demand_present:
        limitations.append("Demand data is unknown because no demand table was mapped; gaps and redeployment are not calculated.")

    return {
        "filename": Path(filename).name[:_SAFE_LABEL_LIMIT],
        "sheet": str(sheet)[:_SAFE_LABEL_LIMIT],
        "supply_mapping": {
            "role_column": supply_role_column[:_SAFE_LABEL_LIMIT],
            "skill_column": supply_skill_column[:_SAFE_LABEL_LIMIT],
            "count_column": supply_count_column[:_SAFE_LABEL_LIMIT] if supply_count_column else None,
            "count_basis": "mapped_column" if supply_count_column else "one_per_row",
        },
        "demand": {
            "status": "known" if demand_present else "unknown",
            "filename": Path(demand_filename).name[:_SAFE_LABEL_LIMIT] if demand_filename else None,
            "sheet": str(demand_sheet)[:_SAFE_LABEL_LIMIT] if demand_sheet else None,
            "mapping": ({
                "role_column": demand_role_column[:_SAFE_LABEL_LIMIT] if demand_role_column else None,
                "skill_column": demand_skill_column[:_SAFE_LABEL_LIMIT] if demand_skill_column else None,
                "count_column": demand_count_column[:_SAFE_LABEL_LIMIT] if demand_count_column else None,
            } if demand_present else None),
        },
        "totals": {
            "supply": _finite(total_supply) if can_show_totals else None,
            "demand": _finite(total_demand) if can_show_totals and demand_present else None,
            "coverage_pct": (
                _finite(total_supply / total_demand * 100)
                if can_show_totals and demand_present and total_demand and total_demand > 0
                else None
            ),
            "gap": _finite(total_gap) if can_show_totals and demand_present else None,
            "surplus": _finite(total_surplus) if can_show_totals and demand_present else None,
            "redeployable_units": _finite(redeployable) if redeployable is not None else None,
        },
        "groups": detail_groups,
        "skill_groups": skill_groups,
        "group_details_withheld": suppressed or totals_suppressed,
        "limitations": limitations,
        "privacy": {
            "mode": "request_scoped_local_processing",
            "persisted": False,
            "source_rows_returned": False,
            "person_level_recommendations": False,
        },
    }


def build_summary_xlsx(profiles: list[dict[str, Any]], language: str = "en") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if language not in {"en", "ar"}:
        raise ValueError("language must be 'en' or 'ar'.")
    is_ar = language == "ar"

    def safe_text(value: Any) -> str:
        text = str(value if value is not None else "")[:32000]
        return "'" + text if text[:1] in {"=", "+", "-", "@"} else text

    workbook = Workbook()
    overview = workbook.active
    overview.title = "نظرة عامة" if is_ar else "Overview"
    overview.append(["THRIVE — ملخص مركز البيانات" if is_ar else "THRIVE Data Hub — Aggregate Summary"])
    overview.append(["تاريخ الإنشاء (UTC)" if is_ar else "Generated (UTC)", pd.Timestamp.now(tz="UTC").isoformat()])
    overview.append([
        "الخصوصية",
        "مخطط البيانات والمؤشرات المجمعة فقط؛ لا يتضمن سجلات المصدر."
        if is_ar else "Schema and aggregate metrics only; no source records included.",
    ])
    overview.append(["الملفات / الأوراق" if is_ar else "Files / sheets", len(profiles)])
    overview.append(["صفوف المصدر" if is_ar else "Source rows", sum(int(item["source_rows"]) for item in profiles)])
    overview.append(["الصفوف المحللة" if is_ar else "Analyzed rows", sum(int(item["analyzed_rows"]) for item in profiles)])

    schema = workbook.create_sheet("ملف البيانات" if is_ar else "Data Profile")
    schema.append(
        ["الملف", "الورقة", "العمود", "النوع", "المفقود", "نسبة المفقود %", "الفريد", "حساس", "المتوسط", "الوسيط", "الحد الأدنى", "الحد الأعلى", "عدد القيم الشاذة"]
        if is_ar else
        ["File", "Sheet", "Column", "Type", "Missing", "Missing %", "Unique", "Sensitive", "Mean", "Median", "Min", "Max", "Outlier count"]
    )
    anomaly_sheet = workbook.create_sheet("إشارات المراجعة" if is_ar else "Anomalies")
    anomaly_sheet.append(
        ["الملف", "الورقة", "الإشارة", "العمود", "المُعلّم", "المقيّم", "النسبة %", "الخطورة", "المنهجية", "الفترات", "سياق الموسمية"]
        if is_ar else
        ["File", "Sheet", "Signal", "Column", "Flagged", "Evaluated", "Rate %", "Severity", "Method", "Periods", "Seasonality context"]
    )
    cleaning_sheet = workbook.create_sheet("سجل التنظيف" if is_ar else "Cleaning Log")
    cleaning_sheet.append(
        ["الملف", "الورقة", "العملية", "الخلايا / الصفوف / الأعمدة المتغيرة"]
        if is_ar else
        ["File", "Sheet", "Operation", "Cells / rows / columns changed"]
    )
    for profile in profiles:
        for column in profile["columns"]:
            stats = column.get("numeric_summary") or {}
            schema.append([
                safe_text(profile["filename"]), safe_text(profile["sheet"]), safe_text(column["name"]),
                safe_text(column["type"]), int(column["missing_count"]), column["missing_percent"],
                int(column["unique_count"]), ("نعم" if column["sensitive"] else "لا") if is_ar else ("Yes" if column["sensitive"] else "No"),
                stats.get("mean"), stats.get("median"), stats.get("min"), stats.get("max"),
                column.get("outlier_count"),
            ])
        for anomaly in profile["anomalies"]:
            anomaly_sheet.append([
                safe_text(profile["filename"]), safe_text(profile["sheet"]),
                safe_text(
                    "ارتفاعات التكلفة الشهرية" if is_ar and anomaly["kind"] == "temporal_cost_spike"
                    else anomaly["kind"]
                ),
                safe_text(anomaly["column"]), int(anomaly["count"]), int(anomaly["evaluated"]),
                anomaly.get("percentage"), safe_text(anomaly["severity"]),
                safe_text(anomaly.get("method_ar") if is_ar else anomaly["method"]),
                safe_text(", ".join(anomaly.get("periods", []))),
                safe_text(anomaly.get("seasonality_context_ar") if is_ar else anomaly.get("seasonality_context", "")),
            ])
        for item in profile["applied_cleaning"]:
            cleaning_sheet.append([
                safe_text(profile["filename"]), safe_text(profile["sheet"]), safe_text(item["id"]), int(item["changed"]),
            ])

    for sheet in workbook.worksheets:
        sheet.sheet_view.rightToLeft = is_ar
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E3070")
        for index, cells in enumerate(sheet.columns, start=1):
            width = min(42, max(12, max((len(str(cell.value or "")) for cell in list(cells)[:80]), default=10) + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_summary_pdf(profiles: list[dict[str, Any]], language: str = "en") -> bytes:
    try:
        from fpdf import FPDF
        from fpdf.enums import XPos, YPos
    except ImportError as exc:
        raise RuntimeError("PDF export is unavailable because fpdf2 is not installed.") from exc

    root = Path(__file__).resolve().parents[2]
    fonts = root / "executive-ui" / "public" / "brand" / "fonts"
    latin_regular = fonts / "font-0.ttf"
    latin_bold = fonts / "font-3.ttf"
    arabic_regular = fonts / "font-4.ttf"
    arabic_bold = fonts / "font-7.ttf"
    if not all(path.is_file() for path in (latin_regular, latin_bold, arabic_regular, arabic_bold)):
        raise RuntimeError("THRIVE report fonts are missing; PDF report cannot be generated.")

    is_ar = language == "ar"
    title = "تقرير مركز البيانات" if is_ar else "THRIVE Data Hub — Executive Summary"
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.add_font("ThriveLatin", style="", fname=str(latin_regular))
    pdf.add_font("ThriveLatin", style="B", fname=str(latin_bold))
    pdf.add_font("ThriveArabic", style="", fname=str(arabic_regular))
    pdf.add_font("ThriveArabic", style="B", fname=str(arabic_bold))

    font = "ThriveArabic" if is_ar else "ThriveLatin"
    def write(text: str, *, bold: bool = False, size: int = 10, height: int = 7) -> None:
        pdf.set_font(font, style="B" if bold else "", size=size)
        if is_ar:
            pdf.set_text_shaping(use_shaping_engine=True, direction="rtl", script="arab", language="ara")
        else:
            pdf.set_text_shaping(use_shaping_engine=False)
        pdf.multi_cell(
            w=pdf.epw,
            h=height,
            text=str(text)[:800],
            align="R" if is_ar else "L",
            new_x=XPos.LMARGIN,
            new_y=YPos.NEXT,
        )

    write(title, bold=True, size=17, height=12)
    write(
        "ملخصات مجمّعة فقط؛ لا يتضمن التقرير سجلات المصدر أو قيماً فردية."
        if is_ar else "Aggregate analysis only; source records and individual values are not included.",
        size=9,
    )
    write(
        f"عدد الملفات / الأوراق: {len(profiles)}"
        if is_ar else f"Files / worksheets: {len(profiles)}",
        bold=True,
    )
    for profile in profiles:
        pdf.ln(3)
        write(f"{profile['filename']} — {profile['sheet']}", bold=True, size=12, height=9)
        if is_ar:
            write(
                f"الصفوف المحللة: {profile['analyzed_rows']} | الأعمدة: {profile['column_count']} | "
                f"جودة البيانات الإرشادية: {profile['quality_score']}/100 | "
                f"إشارات المراجعة: {len(profile['anomalies'])}"
            )
        else:
            write(
                f"Analyzed rows: {profile['analyzed_rows']} | Columns: {profile['column_count']} | "
                f"Heuristic quality: {profile['quality_score']}/100 | Review signals: {len(profile['anomalies'])}"
            )
        for anomaly in profile["anomalies"][:20]:
            anomaly_label = (
                "ارتفاعات التكلفة الشهرية"
                if is_ar and anomaly["kind"] == "temporal_cost_spike"
                else anomaly["kind"]
            )
            write(
                f"• {anomaly_label} — {anomaly['column']}: {anomaly['count']} / "
                f"{anomaly['evaluated']} ({anomaly.get('percentage') or 0}%) — "
                f"{anomaly.get('method_ar') if is_ar else anomaly['method']}",
                size=9,
            )
            periods = ", ".join(anomaly.get("periods", []))
            if periods:
                write(
                    f"الفترات: {periods}" if is_ar else f"Periods: {periods}",
                    size=8,
                )
            seasonal_context = anomaly.get("seasonality_context_ar") if is_ar else anomaly.get("seasonality_context")
            if seasonal_context:
                write(str(seasonal_context), size=8)
        if profile["anomalies"]:
            write("Column-level anomaly counts are review signals, not source records.", size=8)
        for column in profile["columns"][:100]:
            if column["sensitive"]:
                description = f"{column['name']}: sensitive field; value statistics withheld."
            else:
                stats = column.get("numeric_summary") or {}
                description = (
                    f"{column['name']} ({column['type']}): {column['missing_count']} missing, "
                    f"{column['unique_count']} distinct"
                )
                if stats:
                    description += f", median {stats.get('median')}, range {stats.get('min')}–{stats.get('max')}"
            write(description, size=8, height=6)
    if is_ar:
        write("للاستخدام الاستكشافي؛ راجع المنهجية والقيود قبل اتخاذ القرار.", size=8)
    else:
        write("Exploratory analysis only. Review methods and limitations before decisions.", size=8)
    return bytes(pdf.output())
